"""Golden tests for the cover slide.

Locks in the current design against regression. When the Figma design
updates legitimately (and the builder is updated to match), update
these assertions with the new expected values.

Covers (no pun intended):
- Slide dimensions are 1920×1080 (20"×11.25" EMU).
- Exactly one slide is generated.
- The title, subtitle, client, and date strings all appear verbatim.
- The deprecated "Commercial proposal" eyebrow is absent.
- No accent colour is applied to the title.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu


REPO_ROOT = Path(__file__).resolve().parents[2]
SKILL_SCRIPTS = REPO_ROOT / ".claude" / "skills" / "noba-presentation" / "scripts"
SERVICE_ROOT = REPO_ROOT / "service"


@pytest.fixture(scope="module")
def cover_pptx() -> Presentation:
    """Render a known cover spec and return the parsed Presentation."""
    if str(SERVICE_ROOT) not in sys.path:
        sys.path.insert(0, str(SERVICE_ROOT))
    if str(SKILL_SCRIPTS) not in sys.path:
        sys.path.insert(0, str(SKILL_SCRIPTS))

    from models.schemas import DeckSpec, SlideSpec
    from services import deck_builder

    spec = DeckSpec(
        slides=[
            SlideSpec.model_validate({
                "layout": "cover",
                "title": "Next-Generation Coffee Innovation",
                "subtitle": "Identifying and testing new growth opportunities",
                "client": "lavazza",
                "date": "september 2025",
            })
        ]
    )
    data = deck_builder.render(spec)
    return Presentation(io.BytesIO(data))


def _all_text(prs: Presentation) -> str:
    bits: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        bits.append(r.text)
    return " ".join(bits)


def test_single_slide(cover_pptx):
    assert len(cover_pptx.slides) == 1


def test_dimensions_1920x1080(cover_pptx):
    # 1920 px * 9525 EMU/px = 18_288_000 EMU
    assert cover_pptx.slide_width == Emu(18_288_000)
    assert cover_pptx.slide_height == Emu(10_287_000)


def test_title_present(cover_pptx):
    text = _all_text(cover_pptx)
    assert "Next-Generation Coffee Innovation" in text


def test_subtitle_present(cover_pptx):
    assert "Identifying and testing new growth opportunities" in _all_text(cover_pptx)


def test_caption_has_client_and_date(cover_pptx):
    text = _all_text(cover_pptx)
    assert "Prepared for lavazza" in text
    assert "september 2025" in text


def test_no_eyebrow_commercial_proposal(cover_pptx):
    """The old eyebrow was 'Commercial proposal' — must be gone."""
    assert "Commercial proposal" not in _all_text(cover_pptx)


def test_title_is_ink_not_accent(cover_pptx):
    """Title text colour must be ink #1A1A1A. If someone reintroduces
    the accent behaviour this will flag."""
    for slide in cover_pptx.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text == "Next-Generation Coffee Innovation":
                            rgb = r.font.color.rgb
                            assert str(rgb).upper() == "1A1A1A", (
                                f"Title colour is {rgb}, expected 1A1A1A (ink)."
                            )
                            return
    pytest.fail("Could not find the title run to check colour.")


def test_logo_has_N_and_BA(cover_pptx):
    """The wordmark is rendered as two text shapes ('N' and 'BA') plus
    a rounded rectangle pill. Verify at least the two text halves."""
    text = _all_text(cover_pptx)
    # Avoid matching the title "Next-Generation" — check isolated tokens.
    assert " N " in f" {text} " or text.startswith("N ") or text == "N" or " N" in text.split()
    # Look for the "BA" pair surrounded by whitespace or at boundaries.
    tokens = text.split()
    assert "BA" in tokens, f"Expected standalone 'BA' token, got tokens: {tokens}"
