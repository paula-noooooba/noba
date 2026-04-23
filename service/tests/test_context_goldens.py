"""Golden tests for the Context slide.

Locks in the Figma-matched design (node 1:230 of
qqNtw4M8gc3zYRHwKJae7W) against regression.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu


REPO_ROOT = Path(__file__).resolve().parents[2]
SKILL_SCRIPTS = REPO_ROOT / ".claude" / "skills" / "noba-presentation" / "scripts"
SERVICE_ROOT = REPO_ROOT / "service"


SAMPLE_BODY = [
    [
        {"text": "Lavazza aims to expand its vending presence.", "bold": True},
        {"text": " The goal is to elevate the vending experience."},
    ],
    [
        {"text": "BEYOND SIMPLE TRANSACTIONS\n", "bold": True},
        {"text": "Vending is evolving across industries."},
    ],
]


@pytest.fixture(scope="module")
def context_pptx() -> Presentation:
    if str(SERVICE_ROOT) not in sys.path:
        sys.path.insert(0, str(SERVICE_ROOT))
    if str(SKILL_SCRIPTS) not in sys.path:
        sys.path.insert(0, str(SKILL_SCRIPTS))

    from models.schemas import DeckSpec, SlideSpec
    from services import deck_builder

    spec = DeckSpec(
        slides=[SlideSpec.model_validate({
            "layout": "context",
            "tag": "CONTEXT",
            "body": SAMPLE_BODY,
        })]
    )
    data = deck_builder.render(spec)
    return Presentation(io.BytesIO(data))


def _all_text(prs: Presentation) -> str:
    bits: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        bits.append(r.text)
    return " ".join(bits)


def test_single_slide(context_pptx):
    assert len(context_pptx.slides) == 1


def test_dimensions_1920x1080(context_pptx):
    assert context_pptx.slide_width == Emu(18_288_000)
    assert context_pptx.slide_height == Emu(10_287_000)


def test_tag_is_context(context_pptx):
    # Tag comes through UPPERCASED regardless of input case.
    assert "CONTEXT" in _all_text(context_pptx)


def test_body_bold_mixing(context_pptx):
    """The body has both bold and non-bold runs. Verify at least one
    run is bold (the emphasised phrase) and at least one is not."""
    bolds: list[bool] = []
    for slide in context_pptx.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text and ("Lavazza" in r.text or "vending" in r.text):
                            bolds.append(bool(r.font.bold))
    assert True in bolds,  f"No bold run in body; got {bolds}"
    assert False in bolds, f"No regular run in body; got {bolds}"


def test_photo_slot_present(context_pptx):
    """Either a Picture (when spec.image provided) or a rounded-left
    AutoShape (gray placeholder when missing) must occupy the right
    edge of the slide."""
    for slide in context_pptx.slides:
        for shape in slide.shapes:
            # Picture variant
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                return
            # Placeholder variant: AutoShape on the right half
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                if shape.left >= context_pptx.slide_width // 2:
                    return
    pytest.fail("Neither a Picture nor a right-side AutoShape found for the photo slot.")


def test_no_headline_field(context_pptx):
    """Old spec had a big 'headline' at 62px. Figma has no headline —
    make sure we didn't leak one in. No run at that size should exist."""
    for slide in context_pptx.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.font.size.pt >= 45:
                            pytest.fail(
                                f"Found oversized run ({r.font.size.pt}pt) — "
                                f"Context should have no headline: {r.text!r}"
                            )
